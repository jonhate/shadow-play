"""生成参赛文字类材料（作品介绍文档 / 运行说明 / 第三方资源说明 / AI 使用说明）。

输出：D:\\lj\\inktool\\piying\\提交材料\\ 下的 4 个 docx（随后导出 PDF）
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = r"D:\lj\inktool\piying\提交材料"
FONT_CN = "微软雅黑"
FONT_EN = "Segoe UI"


def style_doc(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT_EN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(4)


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = FONT_EN
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        r.font.size = Pt(15 if level == 1 else 12.5)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    return p


def para(doc, text, bold=False, size=10.5, indent=0):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    r = p.add_run(text)
    r.font.name = FONT_EN
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    r.font.size = Pt(size)
    r.bold = bold
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    r.font.name = FONT_EN
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    r.font.size = Pt(size)
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(2)
    return p


def table(doc, rows, widths=None, header=True, size=9.5):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            r = p.add_run(str(val))
            r.font.name = FONT_EN
            r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
            r.font.size = Pt(size)
            r.bold = header and i == 0
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.2
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
    return t


def title_block(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.font.name = FONT_EN
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    r.font.size = Pt(20)
    r.bold = True
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(subtitle)
        r2.font.name = FONT_EN
        r2._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_paragraph()


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    style_doc(doc)
    return doc


os.makedirs(OUT, exist_ok=True)

# ============================================================ 1. 作品介绍文档
doc = new_doc()
title_block(doc, "掌中戲 · 手势皮影", "作品介绍文档")

h(doc, "一、项目定位")
para(doc, "《掌中戲 · 手势皮影》是一个基于普通摄像头的数字皮影戏互动展项。观众站到屏幕前举起双手，"
          "幕布上的皮影便随人而动：张手，皮影的手跟着走；握拳，那只手“执签”定住；两手一合作揖，"
          "双手举高亮相；手掌左右平移，皮影迈步走位。台词不按时间轴自动播完，而要观众做出对应动作才推进一句。")
para(doc, "创作背景：皮影戏是入选联合国教科文组织人类非物质文化遗产代表作名录的传统戏剧，"
          "其魅力在于“耍”而非“看”——艺人一手执三根竹签，一根撑头颈身，两根挑双手，"
          "靠拇指与食指搓动签杆改变关节朝向。但这门手艺需多年练习，展厅里的观众通常只能隔着玻璃观看，"
          "很少有机会亲手操作；而现有手势互动展项多停留在“挥手翻页、隔空点击”，与皮影的操作逻辑无关。")
para(doc, "目标受众与应用场景：博物馆、非遗馆、学校美育课堂与文旅展陈。现场配置为一台普通电脑、"
          "一个 USB 摄像头、一块大屏与音箱；也可通过浏览器在线访问。")
para(doc, "最终成果：可运行的浏览器端互动展项，包含六出戏（三出原创、三出改编）、两套五～六步教学关、"
          "六段戏剧原声背景音乐、三档皮（剪影/镂空/套色）、双人同台、六幕演完的自动谢幕，"
          "以及一键启动、展馆全屏模式与三套自检工具。")

h(doc, "二、主要创新")
table(doc, [
    ["创新点", "原有问题", "采用的方法", "与已有做法的差异", "已取得的实际效果"],
    ["① 把三根竹签的操作逻辑搬到手上",
     "手势互动多为挥手翻页、隔空点击，与皮影操作逻辑无关，观众玩完记不住皮影是什么",
     "张手跟手、握拳执签、两手一合作揖、双手举高亮相、手掌横移走位；进阶模式按传统握法实现"
     "“拇指食指捏合＝夹住竹签、两指连线转动＝搓签转腕”",
     "不是“用手势触发动画”，而是把执签、搓签这两个真实技法做成可操作的动作",
     "六出戏可现场连续操作；两套教学关均可逐步通过（已用自动化脚本逐条验证）"],
    ["② 动作即台词（剧情由手势推进）",
     "台词按时间轴自动播完，观众始终是看客",
     "剧情节拍带 cue 字段，做出对应动作才进入下一句；一次动作只推进一步；"
     "屏幕提示变绿打勾并配锣鼓音效",
     "把“教动作”嵌进“看戏”的过程，观众为看下一句主动学动作",
     "六出戏共 25 个节拍，逐条验证可推进；提示条会在一拍后收起、松手后再出现"],
    ["③ 一只手管一个皮影",
     "单手/双手都控制同一个角色，左右手互相干扰",
     "按画面左右把手分配给两个皮影，每只手只驱动同侧一条手臂；手掌横向位置决定该皮影的走位目标，带死区",
     "把“一人控一偶”的舞台关系做成空间映射，而非手势切换",
     "实测皮影会朝对应那只手的方向走动（如 -318 → -424），回到中间即站住"],
    ["④ 面向展馆的工程化",
     "原型常止步于“能跑”，现场断电、断网、误操作即失效",
     "一键启动脚本与桌面快捷方式、开机自启、全屏 kiosk（藏鼠标/禁右键/出错自愈）、"
     "摄像头与音频自检页、屏内工具面板、实时诊断面板",
     "把非遗展陈的“通电即演、无人值守”当作交付标准",
     "全部资产与依赖本地化，断网可运行；自检页可在不接技术的情况下定位摄像头与音频问题"],
], widths=[3.2, 3.6, 4.2, 3.4, 3.4])

h(doc, "三、技术或创作实现")
para(doc, "技术路线：浏览器端实现，无需安装。渲染使用 three.js 正交相机与逐部件二维仿射矩阵，"
          "按皮影资产（.woop：分件 + 关节 + 竹签）的真实结构逐帧求解姿态；"
          "手部感知使用 MediaPipe Gesture Recognizer（21 个关键点 + 手势分类，单模型）；"
          "声音使用 WebAudio，播放戏剧原声、锣鼓音效与（可选的）台词配音。")
table(doc, [
    ["模块", "职责"],
    ["woop.js", "解析 .woop 资产（zip → 部件贴图 + 骨架数据 + 关节语义）"],
    ["rig.js", "二维仿射、关节链、两段骨骼逆解（纯数学，可离线单测）"],
    ["puppet.js", "单个皮影：网格、逐帧贴矩阵、走步循环、竹签、三档皮肤"],
    ["stage.js", "舞台：正交相机、景片、皮影、自适应布局"],
    ["gesture.js", "MediaPipe 封装：平滑、左右手分配、张开度/捏合/搓动角"],
    ["show.js", "状态机与内容：六出戏、剧情节拍、两套教学关、谢幕时间轴、姿态判定"],
    ["ui.js", "界面：字幕、印章、留影、手部状态牌、教学卡、提示条"],
    ["audio.js", "背景音乐（原声/合成）、锣鼓音效、台词播放与 ducking"],
], widths=[3.0, 13.0])

para(doc, "关键方法（择要）：", bold=True)
bullet(doc, "坐标映射：把舞台坐标换算成骨架坐标时，需按“组原点在头顶”补偿一个身高并翻转符号；"
            "早期版本漏了这一项，导致手的目标点永远落在皮影身下一米多，手臂被顶在极限位，表现为“不跟手、一颤一颤”。")
bullet(doc, "阻尼与死区：关节角以“上一帧”为起点做指数阻尼（而不是从 0 开始），目标点做指数逼近并设死区，"
            "手臂只用到 94% 臂展以避免奇异点，手型判定加滞回，从而抑制关键点噪声。")
bullet(doc, "姿态判定：用双手几何量而非单一手型——两手掌心距离（作揖）、双手高度（亮相）、"
            "两手横向速度同向且够快（甩袖）、两手同时下沉（躬身作揖）。")
bullet(doc, "音频调度：原声解码进内存后循环播放，切幕交叉淡化，台词播放时把背景音乐压到 30%，"
            "并统一走已解锁的 WebAudio 通道（避免浏览器自动播放策略拦截）。")

h(doc, "四、成果状态")
table(doc, [
    ["状态", "内容"],
    ["已完成（可运行、可展示）",
     "六出戏与 25 个剧情节拍；两套教学关（基础 5 步、手指模式 6 步）；三档皮；双人同台；"
     "六幕演完自动谢幕；六段戏剧原声 BGM（裁剪 72 秒、响度归一、淡入淡出、切幕交叉淡化）；"
     "锣鼓动作音效；一键启动/开机自启/展馆全屏/出错自愈；摄像头自检页、音频自检页、实时诊断面板、屏内工具面板；"
     "开源（MIT）并部署在线演示版本。"],
    ["正在完善（已有基础成果，仍在优化）",
     "真机手势阈值需按现场摄像头与光照校准（捏合、作揖、甩袖三处阈值均为可调参数）；"
     "双人同台（4 只手）代码完整，但尚未在两位真人同框的场景下做系统验证。"],
    ["后续计划（尚未实现）",
     "更多剧目与角色扩展；台词配音在“无原声剧目”中的启用；观众留影的扫码带走；"
     "多语言字幕；基于现场实测数据的阈值自适应。"],
], widths=[4.0, 12.0])

h(doc, "五、团队分工与原创贡献")
para(doc, "（按评审匿名化要求，以成员代称）")
table(doc, [
    ["成员", "主要工作", "实际完成成果"],
    ["成员 A", "选题与非遗调研、交互设计、开发与调试、素材授权、测试验收、部署与文档",
     "确定“复刻执签手感”的技术方向；设计张手/执签/作揖/亮相/甩袖与两套教学关；"
     "完成全部功能开发与缺陷修复；取得皮影资产与戏剧原声授权；"
     "编写自动化回归脚本逐项验证；完成本地部署、开源与在线部署"],
    ["AI 编程助手（工具，非成员）", "代码生成与调试辅助、部分文档整理",
     "按成员 A 的设计与反馈生成/修改代码；具体见《AI 工具使用情况说明》"],
    ["第三方开源与素材", "渲染库、手势模型、皮影资产、戏剧原声",
     "见《第三方资源说明》"],
], widths=[3.4, 5.0, 7.6])

h(doc, "六、应用价值")
bullet(doc, "展陈价值：把“看戏”变成“学艺”，观众在 1 分钟内即可完成一次完整的操偶体验；"
            "教学关把非遗“手把手”的传授过程做成了可复制的展陈环节。")
bullet(doc, "传播价值：全离线、单摄像头、普通电脑即可部署，适合基层场馆与学校；"
            "开源（MIT）后可由其他团队替换皮影资产复用整套交互框架。")
bullet(doc, "教育价值：教学关的“必须真做出来才通过”机制，可直接用于非遗进校园的课堂练习。")
para(doc, "说明：本作品尚未开展正式的用户测试与量化评估，以上价值为基于功能实现的判断，"
          "不作为已验证结论表述。", size=9.5)

doc.save(os.path.join(OUT, "03_作品介绍文档", "作品介绍文档.docx") if os.path.isdir(os.path.join(OUT, "03_作品介绍文档"))
         else os.path.join(OUT, "作品介绍文档.docx"))

# ============================================================ 2. 运行或使用说明
doc = new_doc()
title_block(doc, "掌中戲 · 手势皮影", "运行或使用说明")

h(doc, "一、运行环境")
table(doc, [
    ["类别", "要求"],
    ["操作系统", "Windows 10 / 11 64 位（macOS、Linux 同理，只需具备 Node.js 与 Chromium 内核浏览器）"],
    ["运行时", "Node.js 18 及以上（本项目实测 v24.18），仅用于启动本地静态服务器"],
    ["浏览器", "Microsoft Edge 或 Google Chrome（Chromium 100 及以上），需支持 WebGL2、WebAudio、getUserMedia"],
    ["摄像头", "720p 及以上 USB 摄像头（建议 1080p/30fps，程序按 960×540 采集）"],
    ["主机", "近五年普通办公机型即可，集成显卡可运行；建议 i5 八代以上、8GB 内存"],
    ["显示与音响", "1920×1080 及以上；展馆建议 4K 电视或投影 + 普通有源音箱"],
    ["网络", "不需要联网。所有依赖、模型、皮影资产、音频均随包提供"],
], widths=[3.2, 13.2])

h(doc, "二、启动方式（任选其一）")
para(doc, "方式一（推荐，展馆现场）：双击目录下的“启动掌中戏.cmd”，或桌面快捷方式。"
          "程序会自动启动本地服务器（隐藏窗口）并在浏览器打开页面。", bold=False)
para(doc, "方式二（命令行）：在项目目录执行 node tools/serve.mjs，然后在浏览器访问 "
          "http://127.0.0.1:5173/ 。")
para(doc, "方式三（在线演示）：访问已部署的网页版地址（需 HTTPS，浏览器会申请摄像头权限）。")
para(doc, "注意：必须通过 http://127.0.0.1:5173/ 或 https 访问；直接双击 index.html（file://）无法调用摄像头。",
     bold=True)

h(doc, "三、首次使用流程")
table(doc, [
    ["步骤", "操作", "预期结果"],
    ["1", "打开页面", "屏幕上方出现引导卡，右下角小窗显示“▶ 点这里开启摄像头”"],
    ["2", "点击屏幕任意处（或点小窗）", "浏览器弹出“是否允许使用摄像头”，选择“允许”"],
    ["3", "等待小窗出现实时画面", "引导卡切换为动作说明；右侧状态牌显示“左边/右边”"],
    ["4", "按教学关提示逐步做动作", "每步进度条走满即通过，共 5 步（基础模式）"],
    ["5", "教学完成自动开演", "屏幕显示第一出戏与第一句台词"],
], widths=[1.6, 6.4, 8.4])

h(doc, "四、常用操作")
table(doc, [
    ["想做什么", "怎么做"],
    ["开演 / 唤醒", "举一只手，保持约 0.9 秒，屏幕倒数 3-2-1 后开演"],
    ["指挥皮影两只手", "张开手掌移动；手走快了皮影会自己迈步"],
    ["把某只手定住", "握拳（基础模式）；手指模式下松开拇指食指"],
    ["推进剧情", "按屏幕下方金色提示做动作：作揖 / 举高 / 甩袖 / 双拳"],
    ["换下一出戏", "双手举高保持约 2.5 秒（先亮相，再换出）"],
    ["定格留影", "双手握拳保持约 1 秒，落朱印并生成白框照片"],
    ["切换左右手各控一个", "默认开启；两手分别控制画面左右两个皮影，手掌横移带动走位"],
    ["切换手指模式", "点右下角“工具”→手指模式，或按键盘 G"],
    ["重看教学关", "点右上角“教学”，或按键盘 T"],
    ["查看动作总表", "点右上角“？”，或按键盘 H"],
    ["三档皮切换", "按键盘 K（剪影 → 镂空 → 套色）"],
    ["开关背景音乐", "点右下角“音乐”开关，或按键盘 B"],
    ["全屏 / 展馆模式", "“工具”→全屏显示；展馆模式为地址后加 ?kiosk=1"],
    ["双人同台", "“工具”→双人同台（地址后加 ?duo=1），两位观众各站一边"],
], widths=[5.0, 11.4])

h(doc, "五、自检与排障")
table(doc, [
    ["现象", "先看这里", "处理"],
    ["小窗是斜纹红框、没有画面", "框内文字即错误码", "NotAllowedError＝浏览器拒绝：点地址栏摄像头图标允许后，点小窗重试；"
     "NotFoundError＝没找到设备；NotReadableError＝被其它程序占用"],
    ["打开 camtest.html 自检", "权限状态 / 设备列表 / 逐级重试记录", "可复制结果用于排查"],
    ["有画面但皮影不动", "?debug=1 看 hands 是否为 2", "hands 为 0 说明手未被识别：靠近摄像头、提高环境亮度、避免逆光"],
    ["没有声音", "打开 audiotest.html 逐句试听", "看电平条是否跳动；确认页面已被点击过一次（浏览器要求手势后才能出声）"],
    ["动作不灵 / 太灵", "调阈值", "手指模式捏合阈值 show.gripThreshold（默认 0.55）；作揖距离 0.18；甩袖速度 0.8"],
    ["按钮挡住画面元素", "工具→清除记忆后重载", "已在 1080p/1366×768/1280×800 三档验证不重叠"],
], widths=[4.4, 4.6, 7.4])

h(doc, "六、停止与退出")
para(doc, "关闭浏览器不会停止服务器；需要停止时双击“停止掌中戏.cmd”，或关闭后台的 node 进程。"
          "展馆建议把启动脚本放入开机启动项（shell:startup），并使用 ?kiosk=1 全屏模式。")

doc.save(os.path.join(OUT, "运行或使用说明.docx"))

# ============================================================ 3. 第三方资源说明
doc = new_doc()
title_block(doc, "掌中戲 · 手势皮影", "第三方资源说明")
para(doc, "下表列明本项目使用的主要第三方代码、模型、素材及其授权情况。项目自身代码以 MIT 许可发布。")
table(doc, [
    ["资源名称", "类型", "来源", "用途", "授权 / 许可情况"],
    ["three.js (r180)", "开源库", "https://threejs.org", "三维渲染：正交相机、网格与矩阵", "MIT License"],
    ["fflate", "开源库", "https://github.com/101arrowz/fflate", "浏览器端解压 .woop 资产包", "MIT License"],
    ["MediaPipe Tasks-Vision (1.1.1)", "开源库 + 模型", "https://github.com/google-ai-edge/mediapipe",
     "手部 21 关键点与手势分类推理", "Apache-2.0"],
    ["gesture_recognizer.task", "预训练模型", "MediaPipe 官方模型包", "手势识别模型", "随 MediaPipe 发布，Apache-2.0"],
    ["皮影资产 .woop（15 个）", "美术素材", "Woooo Piying Lab（https://piying.woooostudio.com/）",
     "皮影角色与景片", "已获得作者授权用于本项目；第三方使用需自行确认"],
    ["戏剧原声 6 段", "音乐素材", "《牡丹亭·游园》《将军令》《夜深沉》《水龙吟》《万年欢》《良宵引》演奏录音",
     "六出戏背景音乐", "已获授权用于本项目；仅作 72 秒片段裁剪与响度归一，不改变原意"],
    ["台词配音 31 句", "语音素材", "edge-tts（Microsoft 语音合成）于开发期生成", "可选台词配音（当前默认关闭）",
     "按 Microsoft 语音合成服务条款使用；成片不含真人声音"],
    ["界面字体", "字体", "操作系统自带（微软雅黑 / Segoe UI）", "界面与文档排版", "随操作系统授权，未内嵌分发"],
    ["界面图形（回纹、印章、图标）", "原创素材", "本项目自绘 SVG / CSS", "界面装饰", "随项目 MIT 发布"],
], widths=[3.6, 2.2, 4.4, 3.6, 4.2])
para(doc, "说明：")
bullet(doc, "本项目未使用任何需要付费授权的商业字体、图库或音效库；")
bullet(doc, "皮影资产与戏剧原声的授权范围为本次参赛展示与作品演示，如后续商用需另行确认；")
bullet(doc, "开源库均保留其原始许可声明，未修改其许可条款。")

doc.save(os.path.join(OUT, "第三方资源说明.docx"))

# ============================================================ 4. AI 使用情况说明
doc = new_doc()
title_block(doc, "掌中戲 · 手势皮影", "人工智能工具使用情况说明")
para(doc, "一、总体情况")
table(doc, [
    ["项目", "说明"],
    ["是否使用 AI 工具", "是"],
    ["使用的主要工具", "AI 编程助手（代码生成与调试）；edge-tts 语音合成（台词配音）"],
    ["AI 参与环节", "代码实现与调试、台词语音生成、部分文档整理"],
    ["学生自主完成", "选题与非遗调研、交互设计与全部交互语义决策、素材选用与授权、真机测试与阈值调整、"
     "部署与验收、文档定稿"],
], widths=[3.6, 12.8])

para(doc, "二、代表性环节（3 个）", bold=True)
h(doc, "环节 1：皮影骨架驱动与手部映射", level=2)
table(doc, [
    ["项", "内容"],
    ["AI 输入", "“按 .woop 的分件/关节结构，用二维仿射实现逐部件姿态，双手用两段骨骼逆解跟随关键点”"],
    ["AI 输出", "骨架求解、逆解与逐帧贴矩阵的代码初稿"],
    ["学生处理", "用离线渲染脚本核对关节是否错位；发现“手臂只走到目标方向四分之一”与“握拳时手臂掉下来”，"
     "定位到坐标换算漏了身高补偿、以及每帧把关节角清零导致阻尼从 0 开始，修正后手臂才真正跟随"],
    ["最终成果", "皮影手臂可稳定跟随手部目标点，关节角收敛到逆解满值（0.434 → 1.855 弧度）"],
], widths=[2.8, 13.6])

h(doc, "环节 2：手势动作的判定逻辑", level=2)
table(doc, [
    ["项", "内容"],
    ["AI 输入", "“判断作揖、亮相、甩袖、定格四种动作，给出阈值与判定条件”"],
    ["AI 输出", "基于双手几何量的判定函数与阈值"],
    ["学生处理", "真机测试发现“甩袖只成功过一次”“捏着签没法作揖”，反馈现象后修正："
     "甩袖阈值单位错误（把画面宽/秒当成像素/秒）、作揖增加“两手一起下压”这条路径；"
     "并按真机手感调整捏合、作揖、甩袖三处阈值"],
    ["最终成果", "四种动作可稳定触发，且能在教学关中逐步教会观众"],
], widths=[2.8, 13.6])

h(doc, "环节 3：台词语音生成", level=2)
table(doc, [
    ["项", "内容"],
    ["AI 输入", "六出戏的台词文本与角色音色偏好"],
    ["AI 输出", "31 句中文配音 mp3（edge-tts 生成，按出目分配不同音色）"],
    ["学生处理", "决定在已使用戏剧原声的剧目中默认关闭台词配音，避免两层人声叠加；保留开关供无原声剧目使用"],
    ["最终成果", "音频自检页可逐句试听；成片播放以戏剧原声为主"],
], widths=[2.8, 13.6])

para(doc, "三、说明")
bullet(doc, "AI 用于提高实现效率，不参与作品定位、交互语义与素材授权等关键决策；")
bullet(doc, "所有 AI 生成的代码均经学生阅读、修改与实测验证后才纳入作品；")
bullet(doc, "作品中不存在由 AI 生成的、未经核验的功能宣称；")
bullet(doc, "相关 Prompt 与截图等过程材料可按需补充，不在此罗列全部记录。")

doc.save(os.path.join(OUT, "AI工具使用情况说明.docx"))

print("docx 生成完成：", [f for f in os.listdir(OUT) if f.endswith(".docx")])
