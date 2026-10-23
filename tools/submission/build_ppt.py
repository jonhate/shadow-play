"""生成作品汇报 PPT（16:9，非遗红金风格，含真实截图）。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

ROOT = r"D:\lj\inktool\piying\shadow-play"
IMG = os.path.join(ROOT, "docs", "images")
OUT = r"D:\lj\inktool\piying\提交材料\掌中戲_作品汇报PPT.pptx"

W, H = 13.333, 7.5
BG = RGBColor(0x17, 0x10, 0x0A)
BG2 = RGBColor(0x22, 0x16, 0x0C)
GOLD = RGBColor(0xE8, 0xCD, 0x7A)
RED = RGBColor(0xB0, 0x2A, 0x1A)
TEXT = RGBColor(0xF3, 0xE7, 0xD0)
DIM = RGBColor(0xB9, 0xA1, 0x7C)
FONT = "微软雅黑"

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


def set_font(run, size, bold=False, color=TEXT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    ea.set("typeface", FONT)


def bg(slide, color=BG):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def bar(slide, x, y, w, h, color=GOLD):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, lines, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT, spacing=1.25):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if isinstance(line, tuple):
            txt, sz, bd, cl = line
        else:
            txt, sz, bd, cl = line, size, bold, color
        r = p.add_run()
        r.text = txt
        set_font(r, sz, bd, cl)
    return tb


def head(slide, kicker, title, sub=None):
    text(slide, 0.85, 0.5, 11.6, 0.4, [kicker], size=13, color=RED, bold=True)
    text(slide, 0.85, 0.92, 11.6, 0.8, [title], size=32, color=TEXT, bold=True)
    bar(slide, 0.9, 1.78, 1.5, 0.055)
    if sub:
        text(slide, 0.85, 1.95, 11.6, 0.5, [sub], size=14, color=DIM)


def pic(slide, name, x, y, w=None, h=None, caption=None):
    path = os.path.join(IMG, name + ".jpg")
    if not os.path.exists(path):
        return
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)
    if caption:
        text(slide, x, y + p.height / 914400 + 0.06, p.width / 914400, 0.3, [caption], size=11, color=DIM)
    return p


def card(slide, x, y, w, h, title, body, tcolor=GOLD):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = BG2
    s.line.color.rgb = RGBColor(0x4A, 0x35, 0x1C)
    s.line.width = Pt(0.75)
    s.shadow.inherit = False
    text(slide, x + 0.25, y + 0.18, w - 0.5, 0.35, [title], size=15, color=tcolor, bold=True)
    text(slide, x + 0.25, y + 0.62, w - 0.5, h - 0.8, body, size=12.5, color=TEXT, spacing=1.3)


def table(slide, x, y, w, h, rows, col_w=None, size=11.5):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0x2A, 0x1B, 0x0E) if ri == 0 else BG2
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            p = cell.text_frame.paragraphs[0]
            p.line_spacing = 1.15
            r = p.add_run()
            r.text = str(val)
            set_font(r, size, ri == 0, GOLD if ri == 0 else TEXT)
    return tbl


def fill_alpha(shape, opacity):
    """给纯色填充写入真正的透明度（python-pptx 不支持 fill.transparency，
    只能改 XML：<a:srgbClr><a:alpha val="..."/></a:srgbClr>）。
    opacity 1.0 = 不透明，0.0 = 全透明。"""
    spPr = shape._element.spPr
    sf = spPr.find(qn("a:solidFill"))
    srgb = sf.find(qn("a:srgbClr"))
    a = srgb.find(qn("a:alpha"))
    if a is None:
        a = etree.SubElement(srgb, qn("a:alpha"))
    a.set("val", str(int(round(opacity * 100000))))


def pic_cover(slide, path, box_w, box_h):
    """按 cover 方式铺满：保持比例、裁掉溢出部分。"""
    im = Image.open(path)
    iw, ih = im.size
    scale = max(box_w / iw, box_h / ih)
    w, h = iw * scale, ih * scale
    return slide.shapes.add_picture(path, Inches((box_w - w) / 2), Inches((box_h - h) / 2),
                                    width=Inches(w), height=Inches(h))


# ---------------------------------------------------------------- 1 封面
s = prs.slides.add_slide(BLANK)
bg(s)
pic_cover(s, os.path.join(IMG, "cover-scene.png"), W, H)
scrim = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
scrim.fill.solid()
scrim.fill.fore_color.rgb = BG
scrim.line.fill.background()
scrim.shadow.inherit = False
fill_alpha(scrim, 0.26)  # 底图只压一层薄纱，画面清楚可见
left = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(6.5), prs.slide_height)
left.fill.solid()
left.fill.fore_color.rgb = BG
left.line.fill.background()
left.shadow.inherit = False
fill_alpha(left, 0.34)  # 左侧再压一点，保证文字可读
text(s, 0.9, 1.35, 6.0, 0.5, ["2026 数字媒体科技作品"], size=14, color=GOLD, bold=True)
text(s, 0.9, 1.9, 6.0, 1.7, ["掌中戲", "手势皮影"], size=54, color=TEXT, bold=True, spacing=1.1)
bar(s, 0.95, 3.85, 1.8, 0.06)
text(s, 0.9, 4.15, 5.8, 1.2,
     ["举起双手，皮影跟着你动；",
      "做出动作，戏才往下演。"], size=20, color=GOLD, spacing=1.5)
text(s, 0.9, 6.15, 5.9, 0.5, ["浏览器端互动展项 · 无需安装 · 单个摄像头"], size=12, color=DIM)
urlbox = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(6.62), Inches(5.9), Inches(0.52))
urlbox.fill.solid()
urlbox.fill.fore_color.rgb = GOLD
urlbox.line.fill.background()
urlbox.shadow.inherit = False
fill_alpha(urlbox, 0.16)
text(s, 1.05, 6.72, 5.7, 0.4, ["在线演示　jonhate.github.io/shadow-play"], size=13, color=GOLD, bold=True)



# ---------------------------------------------------------------- 2 问题
s = prs.slides.add_slide(BLANK)
bg(s)
head(s, "创作背景", "皮影戏的魅力在“耍”，不在“看”")
card(s, 0.85, 2.5, 3.8, 2.5, "传统是怎么耍的",
     ["艺人一手执三根竹签：一根撑头颈身，两根挑双手；",
      "靠拇指与食指搓动签杆改变关节朝向，",
      "靠抬落平衡杆让皮影进退。"], GOLD)
card(s, 4.85, 2.5, 3.8, 2.5, "难在哪",
     ["手上功夫要练多年；",
      "观众在展厅里只能隔着玻璃看；",
      "非遗“手把手”的传授过程难以复制。"], GOLD)
card(s, 8.85, 2.5, 3.6, 2.5, "现有互动差在哪",
     ["多数手势展项是“挥手翻页、隔空点击”；",
      "交互与皮影的操作逻辑无关；",
      "观众玩完，记不住皮影是什么。"], RED)
text(s, 0.85, 5.45, 11.6, 0.9,
     ["我们要做的不是“用手势控制动画”，而是把三根竹签的操作逻辑直接搬到手上。"],
     size=19, color=GOLD, bold=True)

# ---------------------------------------------------------------- 3 作品是什么
s = prs.slides.add_slide(BLANK)
bg(s)
head(s, "作品定位", "一个能上手“演”的数字皮影戏台")
pic(s, "01-attract", 7.3, 2.15, w=5.2, caption="待机：引导卡 + 点一下开启摄像头")
rows = [
    ["你的手", "皮影的反应"],
    ["张手移动", "对应的那只手跟着你走"],
    ["握拳", "这只手“执签”定住不动"],
    ["两手一合", "作揖 —— 剧情推进一句"],
    ["双手举高", "亮相；举久一点换下一出"],
    ["手掌横移", "皮影迈步走位"],
]
table(s, 0.85, 2.2, 6.1, 3.4, rows, col_w=[2.2, 3.9], size=13)
text(s, 0.85, 5.85, 6.1, 1.0,
     ["六出戏 · 两套教学关 · 六段戏剧原声 · 三档皮 · 双人同台 · 自动谢幕"],
     size=13, color=DIM)

# ---------------------------------------------------------------- 4 创新①
s = prs.slides.add_slide(BLANK)
bg(s)
head(s, "创新 01", "把三根竹签的操作逻辑搬到手上", "不是“手势触发动画”，而是把执签、搓签做成可操作的动作")
pic(s, "Y2-pinch", 7.5, 2.25, w=5.0, caption="手指模式：拇指食指捏住＝夹住竹签")
card(s, 0.85, 2.25, 6.3, 1.55, "基础模式",
     ["张手跟手 · 握拳执签 · 两手一合作揖 · 双手举高亮相 · 手掌横移走位"], GOLD)
card(s, 0.85, 3.95, 6.3, 1.95, "进阶：按传统握法的手指模式",
     ["拇指食指捏住＝夹住竹签，皮影那只手被签带起；",
      "松开＝放签，手臂垂落；",
      "两指连线转动＝传统“搓签”，驱动转腕（以握住那一刻为零点，只算转动量）。"], GOLD)
text(s, 0.85, 6.1, 11.6, 0.6,
     ["依据：传统皮影三根签的分工（平衡杆握法、动作杆在拇指食指间与中指无名指小指间）"],
     size=12, color=DIM)

# ---------------------------------------------------------------- 5 创新②
s = prs.slides.add_slide(BLANK)
bg(s)
head(s, "创新 02", "动作即台词：剧情由手势推进", "台词不按时间轴自动播完，观众从看客变成操偶人")
pic(s, "96-story-cue", 7.5, 2.25, w=5.0, caption="金色提示 → 做出动作 → 变绿打勾 → 推进一句")
card(s, 0.85, 2.25, 6.3, 1.5, "机制",
     ["每句台词带一个 cue（作揖/举高/甩袖/定格）；做出对应动作并保持约 0.45 秒，才进入下一句。"], GOLD)
card(s, 0.85, 3.9, 6.3, 1.5, "防误触",
     ["一次动作只推进一步：不松手不会连跳；提示条推进后收起，松手后再出现。"], GOLD)
card(s, 0.85, 5.55, 6.3, 1.2, "反馈",
     ["提示变绿打勾 + 锣鼓音效；状态牌实时显示“捏住/松开”或“跟手/定住”。"], GOLD)

# ---------------------------------------------------------------- 6 创新③
s = prs.slides.add_slide(BLANK)
bg(s)
head(s, "创新 03", "一只手管一个皮影，手掌横移即走位")
pic(s, "Z3-center", 7.5, 2.25, w=5.0, caption="左手管左边、右手管右边，各自走位")
card(s, 0.85, 2.25, 6.3, 1.6, "空间映射",
     ["按画面左右把手分配给两个皮影；每只手只驱动同侧一条手臂，另一条自然垂落。"], GOLD)
card(s, 0.85, 4.0, 6.3, 1.6, "走位",
     ["皮影的目标位置 = 手掌横向位置 × 0.95，带死区；握拳或松手时站住，到台边自动转身。"], GOLD)
card(s, 0.85, 5.75, 6.3, 1.05, "实测",
     ["双手在画面两侧时，左皮影 −318 → −424、右皮影 +353 → +414，各自朝自己那只手走过去。"], GOLD)

# ---------------------------------------------------------------- 7 技术实现
s = prs.slides.add_slide(BLANK)
bg(s)
head(s, "技术实现", "浏览器端：three.js 渲染 + MediaPipe 手部感知 + WebAudio")
boxes = [
    ("手部感知", "MediaPipe Gesture\nRecognizer\n21 关键点 + 手型", 0.85),
    ("交互与状态", "剧情节拍 / 教学关\n姿态判定 / 模式切换", 3.75),
    ("皮影骨架", "二维仿射 + 两段骨骼逆解\n逐部件贴矩阵", 6.65),
    ("舞台与声音", "正交相机灯箱\n原声 BGM / 锣鼓音效", 9.55),
]
for title, body, x in boxes:
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.2), Inches(2.75), Inches(1.55))
    sh.fill.solid()
    sh.fill.fore_color.rgb = BG2
    sh.line.color.rgb = RGBColor(0x4A, 0x35, 0x1C)
    sh.shadow.inherit = False
    text(s, x + 0.18, 2.35, 2.4, 0.35, [title], size=14, color=GOLD, bold=True)
    text(s, x + 0.18, 2.75, 2.45, 0.95, body.split("\n"), size=11, color=TEXT)
rows = [
    ["模块", "职责"],
    ["woop.js / rig.js", "解析 .woop 资产；仿射、关节链与逆解（纯数学，可离线单测）"],
    ["puppet.js / stage.js", "逐帧姿态、竹签、三档皮肤；正交相机与自适应布局"],
    ["gesture.js", "平滑、左右手分配、张开度/捏合/搓动角"],
    ["show.js", "六出戏与 25 个节拍、两套教学关、谢幕、姿态判定"],
    ["audio.js", "原声 BGM（淡入淡出/交叉淡化）、锣鼓音效、ducking"],
]
table(s, 0.85, 4.0, 11.6, 2.5, rows, col_w=[3.0, 8.6], size=11.5)

# ---------------------------------------------------------------- 8 成果1
s = prs.slides.add_slide(BLANK)
bg(s)
head(s, "实际成果 01", "六出戏，25 个剧情节拍，演完自动谢幕")
pics = [("99-scene-xiang", "西厢 · 月下相会"), ("I4-seg3-walk", "谢幕：换位圆场"), ("21-skin-carved", "三档皮之一：镂空")]
for i, (name, cap) in enumerate(pics):
    pic(s, name, 0.85 + i * 3.95, 2.2, w=3.6, caption=cap)
text(s, 0.85, 5.05, 11.6, 1.9,
     [("六出戏", 13, True, GOLD),
      ("月上柳梢 / 沙场点兵 / 仙翁跨鹤（原创）· 西厢·月下 / 霸王别姬 / 八仙·过海（改编）", 12.5, False, TEXT),
      ("25 个剧情节拍已逐条自动化验证可推进；六幕演完自动演 24 秒谢幕（作揖→亮相→换位→再拜）。", 12.5, False, TEXT)], spacing=1.4)

# ---------------------------------------------------------------- 9 成果2
s = prs.slides.add_slide(BLANK)
bg(s)
head(s, "实际成果 02", "教学关：每一步必须真做出来才通过")
pics = [("94-lesson-bow", "基础教学关第 5/5 步"), ("Y2-pinch", "手指模式教学（捏住＝执签）"), ("18-duo-pose", "双人同台：各控一个皮影")]
for i, (name, cap) in enumerate(pics):
    pic(s, name, 0.85 + i * 3.95, 2.2, w=3.6, caption=cap)
text(s, 0.85, 5.05, 11.6, 1.9,
     [("两套教学关", 13, True, GOLD),
      ("基础模式 5 步（举手 / 移动 / 握拳定住 / 作揖 / 举高）；手指模式 6 步（捏住 / 松开 / 搓动 / 作揖 / 挪动 / 举高）。", 12.5, False, TEXT),
      ("进度按模式分别记忆；三档皮（剪影/镂空/套色）可运行时切换；双人模式 4 只手分组，两位观众各控一个皮影。", 12.5, False, TEXT)], spacing=1.4)

# ---------------------------------------------------------------- 10 工程化
s = prs.slides.add_slide(BLANK)
bg(s)
head(s, "工程化", "面向展馆：通电即演、断网可用、出错自愈")
pic(s, "T3-debug-on", 7.3, 2.2, w=5.2, caption="屏内工具面板 + 实时诊断")
rows = [
    ["能力", "做法"],
    ["一键启动", "双击启动脚本：自动起本地服务器并打开浏览器"],
    ["开机自启", "快捷方式放入 shell:startup"],
    ["展馆全屏", "?kiosk=1：隐藏鼠标、屏蔽右键、出错自动重载"],
    ["离线运行", "渲染库、手势模型、皮影资产、音频全部本地化"],
    ["自检工具", "摄像头自检页、音频自检页、?debug=1 实时面板、屏内工具面板"],
    ["可回归", "一键脚本验证 25 个节拍 / 两套教学关 / 谢幕 / 三档皮 / 六种模式"],
]
table(s, 0.85, 2.2, 6.2, 3.6, rows, col_w=[1.9, 4.3], size=11.5)
text(s, 0.85, 5.95, 6.2, 1.1,
     [("单文件部署：普通办公机 + 一个 USB 摄像头即可", 12.5, False, DIM),
      ("在线演示：jonhate.github.io/shadow-play", 12.5, True, GOLD)], spacing=1.5)

# ---------------------------------------------------------------- 11 原创贡献与 AI
s = prs.slides.add_slide(BLANK)
bg(s)
head(s, "原创贡献", "学生做了什么，AI 做了什么（如实说明）")
rows = [
    ["环节", "成员 A（参赛学生）", "AI 编程助手"],
    ["选题与调研", "非遗调研、对比现有手势展项，确定“复刻执签手感”方向", "—"],
    ["交互设计", "动作语义（握拳=执签、作揖=剧情闸门）、两套教学关、提示与反馈", "—"],
    ["代码实现", "提出方案、逐项实测、定位并修复缺陷、验收", "按设计生成/修改代码初稿"],
    ["素材与授权", "皮影资产、六段戏剧原声的选用与授权", "—"],
    ["测试与部署", "真机调阈值、自动化回归脚本、本地与在线部署", "—"],
    ["台词配音", "决定有原声时默认关闭配音", "edge-tts 生成 31 句配音"],
]
table(s, 0.85, 2.15, 11.6, 3.5, rows, col_w=[2.0, 5.6, 4.0], size=11)
text(s, 0.85, 5.9, 11.6, 1.1,
     ["AI 用于提高实现效率，不参与作品定位、交互语义与素材授权等关键决策；所有 AI 生成代码均经阅读、修改与实测后才纳入作品。"],
     size=12.5, color=GOLD)

# ---------------------------------------------------------------- 12 价值
s = prs.slides.add_slide(BLANK)
bg(s)
head(s, "价值", "把“看戏”变成“学艺”")
card(s, 0.85, 2.2, 3.75, 2.5, "展陈",
     ["观众 1 分钟内可完成一次完整操偶；",
      "教学关把非遗“手把手”的传授过程做成可复制的展陈环节。"], GOLD)
card(s, 4.8, 2.2, 3.75, 2.5, "传播",
     ["单摄像头、普通电脑、断网可用；",
      "适合基层场馆与学校；开源（MIT）后可替换皮影资产复用交互框架。"], GOLD)
card(s, 8.75, 2.2, 3.7, 2.5, "教育",
     ["“必须真做出来才通过”的机制，",
      "可直接用于非遗进校园的课堂练习。"], GOLD)
text(s, 0.85, 5.1, 11.6, 1.4,
     [("边界说明", 12.5, True, RED),
      ("本作品尚未开展正式用户测试与量化评估；真机手势阈值需按现场摄像头与光照校准；双人 4 手模式代码完整但未在真人双人场景下系统验证。", 11.5, False, DIM)], spacing=1.4)

# ---------------------------------------------------------------- 13 结束
s = prs.slides.add_slide(BLANK)
bg(s)
pic_cover(s, os.path.join(IMG, "cover-scene.png"), W, H)
ov2 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
ov2.fill.solid()
ov2.fill.fore_color.rgb = BG
ov2.line.fill.background()
ov2.shadow.inherit = False
fill_alpha(ov2, 0.34)
text(s, 0, 2.35, W, 1.0, ["掌中戲 · 手势皮影"], size=44, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
bar(s, W / 2 - 1.1, 3.5, 2.2, 0.06)
text(s, 0, 3.8, W, 0.6, ["一盏灯，一张皮，一双手"], size=20, color=GOLD, align=PP_ALIGN.CENTER)
qbox = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(W / 2 - 3.4), Inches(4.85), Inches(6.8), Inches(0.95))
qbox.fill.solid()
qbox.fill.fore_color.rgb = GOLD
qbox.line.fill.background()
qbox.shadow.inherit = False
fill_alpha(qbox, 0.16)
text(s, 0, 4.95, W, 0.4, ["在线试玩"], size=12, color=DIM, align=PP_ALIGN.CENTER)
text(s, 0, 5.25, W, 0.5, ["https://jonhate.github.io/shadow-play/"], size=19, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
text(s, 0, 6.25, W, 0.5, ["浏览器打开 → 点击页面并允许摄像头 → 跟随教学关即可上手；首次加载约 10–20 秒"],
     size=12, color=DIM, align=PP_ALIGN.CENTER)

prs.save(OUT)
print("PPT 已生成：", OUT, "共", len(prs.slides.__iter__.__self__._sldIdLst), "页")
